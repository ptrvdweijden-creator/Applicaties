"""Douchebeslag voorraadbeheer API.

Datamodel (uitbreidbaar):
- users: gebruikers met rol (admin / magazijn / monteur)
- articles: artikelstam; artikelnummer uniek en onwijzigbaar; barcodes uniek over alle artikelen
- mutations: onveranderlijk grootboek van voorraadmutaties (nooit overschreven/verwijderd)
- projects / reservations: voorbereid voor fase 2; gereserveerde voorraad = som actieve reserveringen
"""
import io
import os
import re
import logging
import uuid
import xml.etree.ElementTree as ET
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Annotated, Any, List, Literal, Optional

import jwt
import requests
from bson import ObjectId
from dotenv import load_dotenv
from fastapi import APIRouter, Depends, FastAPI, File, HTTPException, Query, UploadFile
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from jwt.exceptions import InvalidTokenError
from motor.motor_asyncio import AsyncIOMotorClient
from pwdlib import PasswordHash
from pydantic import BaseModel, BeforeValidator, ConfigDict, Field, field_validator
from pymongo import ReturnDocument
from pymongo.errors import DuplicateKeyError
from starlette.middleware.cors import CORSMiddleware

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")

client = AsyncIOMotorClient(os.environ["MONGO_URL"], tz_aware=True)
db = client[os.environ["DB_NAME"]]
JWT_SECRET = os.environ["JWT_SECRET"]
TOKEN_HOURS = int(os.environ["ACCESS_TOKEN_HOURS"])

# ---------------------------------------------------------------- Object storage (Emergent managed)
STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
EMERGENT_KEY = os.environ["EMERGENT_LLM_KEY"]
STORAGE_APP_NAME = os.environ["STORAGE_APP_NAME"]
storage_key: Optional[str] = None


def init_storage() -> str:
    global storage_key
    if storage_key:
        return storage_key
    resp = requests.post(f"{STORAGE_URL}/init", json={"emergent_key": EMERGENT_KEY}, timeout=30)
    resp.raise_for_status()
    storage_key = resp.json()["storage_key"]
    return storage_key


def _storage_call(method: str, path: str, **kw) -> requests.Response:
    global storage_key
    for attempt in range(2):
        resp = requests.request(method, f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": init_storage(), **kw.pop("headers", {})}, timeout=60, **kw)
        if resp.status_code == 503 and attempt == 0:
            storage_key = None
            continue
        break
    if resp.status_code == 402:
        raise HTTPException(402, "Opslagtegoed is op. Uploaden is tijdelijk niet mogelijk.")
    resp.raise_for_status()
    return resp


def put_object(path: str, data: bytes, content_type: str) -> dict:
    return _storage_call("PUT", path, data=data, headers={"Content-Type": content_type}).json()


def get_object(path: str) -> bytes:
    return _storage_call("GET", path).content


hasher = PasswordHash.recommended()
bearer = HTTPBearer(auto_error=False)
logger = logging.getLogger("voorraad")
logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(name)s - %(levelname)s - %(message)s")

app = FastAPI(title="Douchebeslag Voorraad API")
api = APIRouter(prefix="/api")

Role = Literal["admin", "magazijn", "monteur"]
MutationType = Literal["ontvangst", "afboeken", "retour", "correctie"]

# Rechten per mutatietype
MUTATION_ROLES = {
    "afboeken": {"admin", "magazijn", "monteur"},
    "retour": {"admin", "magazijn", "monteur"},
    "ontvangst": {"admin", "magazijn"},
    "correctie": {"admin", "magazijn"},
}


def now() -> datetime:
    return datetime.now(timezone.utc)


PyObjectId = Annotated[str, BeforeValidator(lambda v: str(v) if isinstance(v, ObjectId) else v)]


class BaseDocument(BaseModel):
    model_config = ConfigDict(populate_by_name=True)
    id: Optional[PyObjectId] = Field(default=None, alias="_id")

    @classmethod
    def from_mongo(cls, doc: dict):
        return cls.model_validate(doc)

    def to_mongo(self) -> dict:
        data = self.model_dump(by_alias=True, exclude={"id"})
        if self.id:
            data["_id"] = ObjectId(self.id)
        return data


def oid(value: str) -> ObjectId:
    try:
        return ObjectId(value)
    except Exception:
        raise HTTPException(404, "Niet gevonden")


# ---------------------------------------------------------------- Models
class User(BaseDocument):
    email: str
    name: str
    role: Role
    active: bool = True
    password_hash: str = ""
    created_at: datetime = Field(default_factory=now)

    def public(self) -> dict:
        return {"id": self.id, "email": self.email, "name": self.name, "role": self.role, "active": self.active}


class Article(BaseDocument):
    article_number: str
    name: str = ""
    description: str
    category: str
    barcodes: List[str] = []
    unit: str = "stuks"
    physical_stock: int = 0
    min_stock: int = 0
    notes: str = ""
    active: bool = True
    # Visuele barcode (SVG in object storage) + de daadwerkelijk gecodeerde waarde (moet in barcodes staan)
    barcode_svg: Optional[dict] = None
    import_batch: Optional[str] = None
    created_at: datetime = Field(default_factory=now)
    updated_at: datetime = Field(default_factory=now)

    @property
    def title(self) -> str:
        return self.name or self.description


class Mutation(BaseDocument):
    client_request_id: str
    article_id: PyObjectId
    article_number: str
    article_description: str
    type: Literal["ontvangst", "afboeken", "retour", "correctie", "beginvoorraad"]
    quantity: int  # ingevoerd aantal (bij correctie: getelde voorraad)
    delta: int = 0  # werkelijke wijziging van fysieke voorraad
    stock_before: Optional[int] = None
    stock_after: Optional[int] = None
    status: Literal["in_behandeling", "geboekt", "afgewezen"] = "in_behandeling"
    reject_reason: Optional[str] = None
    user_id: PyObjectId
    user_name: str
    note: str = ""
    project_id: Optional[PyObjectId] = None  # toekomstige koppeling aan project
    created_at: datetime = Field(default_factory=now)


# Requests
class LoginIn(BaseModel):
    email: str
    password: str


class UserCreateIn(BaseModel):
    email: str
    name: str = Field(min_length=1, max_length=100)
    role: Role
    password: str = Field(min_length=8)

    @field_validator("email")
    @classmethod
    def _email(cls, v: str) -> str:
        v = v.strip().lower()
        if not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", v):
            raise ValueError("Ongeldig e-mailadres")
        return v


class UserUpdateIn(BaseModel):
    name: Optional[str] = None
    role: Optional[Role] = None
    active: Optional[bool] = None
    password: Optional[str] = Field(default=None, min_length=8)


def _clean_barcodes(v: List[str]) -> List[str]:
    out = []
    for b in v:
        b = b.strip()
        if b and b not in out:
            out.append(b)
    return out


class ArticleCreateIn(BaseModel):
    article_number: str = Field(min_length=1, max_length=50)
    name: str = Field(default="", max_length=200)
    description: str = Field(min_length=1, max_length=200)
    category: str = Field(min_length=1, max_length=60)
    barcodes: List[str] = []
    unit: str = Field(default="stuks", min_length=1, max_length=20)
    min_stock: int = Field(default=0, ge=0)
    initial_stock: int = Field(default=0, ge=0)
    notes: str = ""
    active: bool = True

    @field_validator("barcodes")
    @classmethod
    def _bc(cls, v):
        return _clean_barcodes(v)

    @field_validator("article_number", "description", "category", "unit")
    @classmethod
    def _strip(cls, v: str) -> str:
        v = v.strip()
        if not v:
            raise ValueError("Verplicht veld")
        return v


class ArticleUpdateIn(BaseModel):
    model_config = ConfigDict(extra="forbid")  # artikelnummer en voorraad zijn hier niet wijzigbaar
    name: Optional[str] = Field(default=None, max_length=200)
    description: Optional[str] = Field(default=None, min_length=1, max_length=200)
    category: Optional[str] = Field(default=None, min_length=1, max_length=60)
    barcodes: Optional[List[str]] = None
    unit: Optional[str] = Field(default=None, min_length=1, max_length=20)
    min_stock: Optional[int] = Field(default=None, ge=0)
    notes: Optional[str] = None
    active: Optional[bool] = None

    @field_validator("barcodes")
    @classmethod
    def _bc(cls, v):
        return _clean_barcodes(v) if v is not None else v


class BarcodeLinkIn(BaseModel):
    barcode: str = Field(min_length=1, max_length=100)


class MutationIn(BaseModel):
    client_request_id: str = Field(min_length=8, max_length=80)
    article_id: str
    type: MutationType
    quantity: int = Field(ge=0, le=100000)
    note: str = Field(default="", max_length=500)


# ---------------------------------------------------------------- Auth
def make_token(user: User) -> str:
    t = now()
    return jwt.encode({"sub": user.id, "role": user.role, "iat": t, "exp": t + timedelta(hours=TOKEN_HOURS)}, JWT_SECRET, algorithm="HS256")


async def current_user(cred: Optional[HTTPAuthorizationCredentials] = Depends(bearer)) -> User:
    err = HTTPException(401, "Sessie verlopen of ongeldig. Log opnieuw in.")
    if not cred:
        raise err
    try:
        payload = jwt.decode(cred.credentials, JWT_SECRET, algorithms=["HS256"])
        doc = await db.users.find_one({"_id": ObjectId(payload["sub"])})
    except (InvalidTokenError, KeyError, Exception):
        raise err
    if not doc or not doc.get("active", True):
        raise err
    return User.from_mongo(doc)


def require_roles(*roles: str):
    async def dep(user: User = Depends(current_user)) -> User:
        if user.role not in roles:
            raise HTTPException(403, "Onvoldoende rechten voor deze actie")
        return user
    return dep


@api.post("/auth/login")
async def login(data: LoginIn):
    doc = await db.users.find_one({"email": data.email.strip().lower()})
    user = User.from_mongo(doc) if doc else None
    if not user or not hasher.verify(data.password, user.password_hash) or not user.active:
        raise HTTPException(401, "E-mailadres of wachtwoord is onjuist")
    return {"access_token": make_token(user), "user": user.public()}


@api.get("/auth/me")
async def me(user: User = Depends(current_user)):
    return user.public()


# ---------------------------------------------------------------- Users (admin)
@api.get("/users")
async def list_users(_: User = Depends(require_roles("admin"))):
    docs = await db.users.find().sort("name", 1).to_list(500)
    return [User.from_mongo(d).public() for d in docs]


@api.post("/users", status_code=201)
async def create_user(data: UserCreateIn, _: User = Depends(require_roles("admin"))):
    user = User(email=data.email, name=data.name.strip(), role=data.role, password_hash=hasher.hash(data.password))
    try:
        res = await db.users.insert_one(user.to_mongo())
    except DuplicateKeyError:
        raise HTTPException(409, "Dit e-mailadres is al in gebruik")
    user.id = str(res.inserted_id)
    return user.public()


@api.patch("/users/{user_id}")
async def update_user(user_id: str, data: UserUpdateIn, admin: User = Depends(require_roles("admin"))):
    upd: dict[str, Any] = {}
    if data.name is not None:
        upd["name"] = data.name.strip()
    if data.role is not None:
        upd["role"] = data.role
    if data.active is not None:
        upd["active"] = data.active
    if data.password:
        upd["password_hash"] = hasher.hash(data.password)
    if user_id == admin.id and (upd.get("active") is False or upd.get("role", "admin") != "admin"):
        raise HTTPException(400, "U kunt uw eigen beheerdersaccount niet deactiveren of degraderen")
    doc = await db.users.find_one_and_update({"_id": oid(user_id)}, {"$set": upd}, return_document=ReturnDocument.AFTER)
    if not doc:
        raise HTTPException(404, "Gebruiker niet gevonden")
    return User.from_mongo(doc).public()


# ---------------------------------------------------------------- Articles
async def reserved_map(article_ids: List[ObjectId]) -> dict:
    """Gereserveerde voorraad per artikel (actieve reserveringen). Fase 2: projectreserveringen."""
    rows = await db.reservations.aggregate([
        {"$match": {"article_id": {"$in": article_ids}, "status": "actief"}},
        {"$group": {"_id": "$article_id", "qty": {"$sum": "$quantity"}}},
    ]).to_list(None)
    return {str(r["_id"]): r["qty"] for r in rows}


def article_out(a: Article, reserved: int = 0) -> dict:
    d = a.model_dump()
    d.pop("created_at", None)
    d.pop("import_batch", None)
    if a.barcode_svg:
        svg = a.barcode_svg
        d["barcode_svg"] = {"filename": svg["filename"], "barcode_value": svg["barcode_value"], "size": svg["size"],
                            "version": svg["path"], "uploaded_at": svg["uploaded_at"].isoformat()}
    d["updated_at"] = a.updated_at.isoformat()
    d["reserved_stock"] = reserved
    d["available_stock"] = a.physical_stock - reserved
    d["below_min"] = a.physical_stock < a.min_stock
    return d


async def articles_out(docs: List[dict]) -> List[dict]:
    arts = [Article.from_mongo(d) for d in docs]
    res = await reserved_map([ObjectId(a.id) for a in arts])
    return [article_out(a, res.get(a.id, 0)) for a in arts]


async def ensure_barcodes_free(barcodes: List[str], exclude_id: Optional[str] = None):
    if not barcodes:
        return
    q: dict[str, Any] = {"barcodes": {"$in": barcodes}}
    if exclude_id:
        q["_id"] = {"$ne": ObjectId(exclude_id)}
    clash = await db.articles.find_one(q)
    if clash:
        used = [b for b in barcodes if b in clash["barcodes"]]
        raise HTTPException(409, f"Barcode {', '.join(used)} is al gekoppeld aan artikel {clash['article_number']}")


@api.get("/categories")
async def categories(_: User = Depends(current_user)):
    return sorted(await db.articles.distinct("category"))


@api.get("/articles")
async def list_articles(
    q: str = "",
    category: str = "",
    status: Literal["actief", "inactief", "alle"] = "actief",
    below_min: bool = False,
    limit: int = Query(500, le=1000),
    _: User = Depends(current_user),
):
    filt: dict[str, Any] = {}
    if status != "alle":
        filt["active"] = status == "actief"
    if category:
        filt["category"] = category
    if q.strip():
        rx = {"$regex": re.escape(q.strip()), "$options": "i"}
        filt["$or"] = [{"article_number": rx}, {"name": rx}, {"description": rx}, {"barcodes": rx}]
    if below_min:
        filt["$expr"] = {"$lt": ["$physical_stock", "$min_stock"]}
    docs = await db.articles.find(filt).sort("article_number", 1).to_list(limit)
    return await articles_out(docs)


@api.get("/articles/by-barcode/{code}")
async def by_barcode(code: str, _: User = Depends(current_user)):
    doc = await db.articles.find_one({"barcodes": code.strip()})
    if not doc:
        raise HTTPException(404, "Onbekende barcode")
    return (await articles_out([doc]))[0]


@api.get("/articles/{article_id}")
async def get_article(article_id: str, _: User = Depends(current_user)):
    doc = await db.articles.find_one({"_id": oid(article_id)})
    if not doc:
        raise HTTPException(404, "Artikel niet gevonden")
    return (await articles_out([doc]))[0]


@api.post("/articles", status_code=201)
async def create_article(data: ArticleCreateIn, user: User = Depends(require_roles("admin"))):
    await ensure_barcodes_free(data.barcodes)
    art = Article(**data.model_dump(exclude={"initial_stock"}), physical_stock=0)
    try:
        res = await db.articles.insert_one(art.to_mongo())
    except DuplicateKeyError as e:
        if "article_number" in str(e):
            raise HTTPException(409, f"Artikelnummer {data.article_number} bestaat al")
        raise HTTPException(409, "Een van de barcodes is al in gebruik")
    art.id = str(res.inserted_id)
    if data.initial_stock > 0:
        await apply_initial_stock(art, data.initial_stock, user, "Beginvoorraad bij aanmaken artikel")
    return await get_article(art.id, user)


async def apply_initial_stock(art: Article, qty: int, user: User, note: str):
    """Beginvoorraad: één keer vastgelegd als 'beginvoorraad'-mutatie (geen ontvangst), alleen als de stand nog 0 is."""
    mut = Mutation(client_request_id=f"init-{art.id}", article_id=art.id, article_number=art.article_number,
                   article_description=art.title, type="beginvoorraad", quantity=qty, delta=qty, stock_before=0,
                   stock_after=qty, user_id=user.id, user_name=user.name, note=note)
    try:
        res = await db.mutations.insert_one(mut.to_mongo())
    except DuplicateKeyError:
        return  # beginvoorraad al geboekt
    upd = await db.articles.update_one({"_id": ObjectId(art.id), "physical_stock": 0},
                                       {"$set": {"physical_stock": qty, "updated_at": now()}})
    status = {"status": "geboekt"} if upd.modified_count else {"status": "afgewezen", "reject_reason": "Voorraad was niet 0"}
    await db.mutations.update_one({"_id": res.inserted_id}, {"$set": status})


@api.patch("/articles/{article_id}")
async def update_article(article_id: str, data: ArticleUpdateIn, user: User = Depends(require_roles("admin"))):
    upd = {k: v for k, v in data.model_dump().items() if v is not None}
    for k in ("description", "category", "unit"):
        if k in upd:
            upd[k] = upd[k].strip()
    if "barcodes" in upd:
        await ensure_barcodes_free(upd["barcodes"], article_id)
        cur = await db.articles.find_one({"_id": oid(article_id)}, {"barcode_svg": 1})
        svg = (cur or {}).get("barcode_svg")
        if svg and svg["barcode_value"] not in upd["barcodes"]:
            raise HTTPException(409, f"Barcode {svg['barcode_value']} hoort bij de gekoppelde SVG. Verwijder eerst de SVG-barcode.")
    upd["updated_at"] = now()
    try:
        doc = await db.articles.find_one_and_update({"_id": oid(article_id)}, {"$set": upd}, return_document=ReturnDocument.AFTER)
    except DuplicateKeyError:
        raise HTTPException(409, "Een van de barcodes is al in gebruik")
    if not doc:
        raise HTTPException(404, "Artikel niet gevonden")
    return (await articles_out([doc]))[0]


@api.post("/articles/{article_id}/barcodes")
async def link_barcode(article_id: str, data: BarcodeLinkIn, user: User = Depends(require_roles("admin"))):
    code = data.barcode.strip()
    await ensure_barcodes_free([code], article_id)
    try:
        doc = await db.articles.find_one_and_update(
            {"_id": oid(article_id)}, {"$addToSet": {"barcodes": code}, "$set": {"updated_at": now()}},
            return_document=ReturnDocument.AFTER)
    except DuplicateKeyError:
        raise HTTPException(409, "Deze barcode is al in gebruik")
    if not doc:
        raise HTTPException(404, "Artikel niet gevonden")
    return (await articles_out([doc]))[0]


# ---------------------------------------------------------------- Excel-import artikelen
IMPORT_COLUMNS = ["Artikelnaam", "Artikelcode", "Artikelbeschrijving", "Hoeveelheid", "Categorie", "Minimumhoeveelheid"]
IMPORT_REQUIRED = {"Artikelnaam", "Artikelcode", "Hoeveelheid", "Categorie", "Minimumhoeveelheid"}
MAX_IMPORT_ROWS = 2000


def _norm_header(v: Any) -> str:
    return re.sub(r"\s+", "", str(v or "")).lower()


def _cell_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)  # Excel slaat codes soms als getal op (12345.0)
    return str(v).strip()


def _cell_int(v: Any, label: str, errors: List[str]) -> Optional[int]:
    t = _cell_text(v).replace(",", ".")
    if t == "":
        errors.append(f"{label} ontbreekt")
        return None
    try:
        f = float(t)
    except ValueError:
        errors.append(f"{label} '{t}' is geen getal")
        return None
    if not f.is_integer():
        errors.append(f"{label} moet een geheel getal zijn")
        return None
    if f < 0:
        errors.append(f"{label} mag niet negatief zijn")
        return None
    return int(f)


async def validate_import_rows(raw: List[dict]) -> List[dict]:
    """Valideert regels (uit Excel of uit de bevestigde preview). Controleert dubbele codes in bestand én database."""
    existing_codes = {c.lower() for c in await db.articles.distinct("article_number")}
    categories = {c.lower(): c for c in await db.articles.distinct("category")}
    seen: dict[str, int] = {}
    out = []
    for r in raw:
        errors: List[str] = []
        code = _cell_text(r.get("Artikelcode"))
        name = _cell_text(r.get("Artikelnaam"))
        desc = _cell_text(r.get("Artikelbeschrijving"))
        cat = _cell_text(r.get("Categorie"))
        if not code:
            errors.append("Artikelcode ontbreekt")
        elif len(code) > 50:
            errors.append("Artikelcode is langer dan 50 tekens")
        if not name:
            errors.append("Artikelnaam ontbreekt")
        elif len(name) > 200:
            errors.append("Artikelnaam is langer dan 200 tekens")
        if len(desc) > 200:
            errors.append("Artikelbeschrijving is langer dan 200 tekens")
        if not cat:
            errors.append("Categorie ontbreekt")
        elif len(cat) > 60:
            errors.append("Categorie is langer dan 60 tekens")
        qty = _cell_int(r.get("Hoeveelheid"), "Hoeveelheid", errors)
        mn = _cell_int(r.get("Minimumhoeveelheid"), "Minimumhoeveelheid", errors)
        if code:
            key = code.lower()
            if key in existing_codes:
                errors.append("Artikelcode bestaat al in de database (wordt niet overschreven)")
            if key in seen:
                errors.append(f"Dubbele artikelcode in bestand (ook op regel {seen[key]})")
            else:
                seen[key] = r["row"]
        if cat:
            cat = categories.setdefault(cat.lower(), cat)  # zelfde schrijfwijze als bestaande categorie
        out.append({"row": r["row"], "Artikelcode": code, "Artikelnaam": name, "Artikelbeschrijving": desc,
                    "Categorie": cat, "Hoeveelheid": qty, "Minimumhoeveelheid": mn, "errors": errors, "valid": not errors})
    return out


@api.get("/import/articles/template")
async def import_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Artikelen"
    ws.append(IMPORT_COLUMNS)
    for i, col in enumerate(IMPORT_COLUMNS):
        ws.cell(row=1, column=i + 1).font = Font(bold=True)
        ws.column_dimensions[chr(65 + i)].width = 24
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="voorbeeld-artikelimport.xlsx"'})


@api.post("/import/articles/preview")
async def import_preview(file: UploadFile = File(...), _: User = Depends(require_roles("admin"))):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(400, "Alleen Excelbestanden (.xlsx) worden ondersteund")
    content = await file.read()
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(400, "Bestand is groter dan 5 MB")
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
    except Exception:
        raise HTTPException(400, "Bestand kon niet worden gelezen. Upload een geldig .xlsx-bestand.")
    header_idx = next((i for i, r in enumerate(rows) if any(_cell_text(v) for v in r)), None)
    if header_idx is None:
        raise HTTPException(400, "Het bestand is leeg")
    header = {_norm_header(v): i for i, v in enumerate(rows[header_idx]) if _cell_text(v)}
    missing = [c for c in IMPORT_COLUMNS if _norm_header(c) not in header and c in IMPORT_REQUIRED]
    if missing:
        raise HTTPException(400, f"Ontbrekende kolom(men): {', '.join(missing)}. Gebruik het voorbeeldbestand.")
    raw = []
    for i, r in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        vals = {c: (r[header[_norm_header(c)]] if _norm_header(c) in header and header[_norm_header(c)] < len(r) else None) for c in IMPORT_COLUMNS}
        if not any(_cell_text(v) for v in vals.values()):
            continue  # lege regel
        raw.append({"row": i, **vals})
    if not raw:
        raise HTTPException(400, "Het bestand bevat geen artikelregels")
    if len(raw) > MAX_IMPORT_ROWS:
        raise HTTPException(400, f"Maximaal {MAX_IMPORT_ROWS} regels per import")
    result = await validate_import_rows(raw)
    return {"filename": file.filename, "rows": result, "valid_count": sum(r["valid"] for r in result),
            "invalid_count": sum(not r["valid"] for r in result)}


class ImportRowIn(BaseModel):
    row: int
    Artikelcode: Any = None
    Artikelnaam: Any = None
    Artikelbeschrijving: Any = None
    Hoeveelheid: Any = None
    Categorie: Any = None
    Minimumhoeveelheid: Any = None


class ImportCommitIn(BaseModel):
    filename: str = Field(default="", max_length=200)
    rows: List[ImportRowIn] = Field(max_length=MAX_IMPORT_ROWS)


@api.post("/import/articles/commit")
async def import_commit(data: ImportCommitIn, user: User = Depends(require_roles("admin"))):
    checked = await validate_import_rows([r.model_dump() for r in data.rows])  # opnieuw controleren vlak voor opslaan
    batch = f"import-{now().strftime('%Y%m%d%H%M%S')}-{uuid.uuid4().hex[:6]}"
    added, skipped = [], []
    for r in checked:
        if not r["valid"]:
            skipped.append({"row": r["row"], "article_number": r["Artikelcode"], "reason": "; ".join(r["errors"])})
            continue
        art = Article(article_number=r["Artikelcode"], name=r["Artikelnaam"], description=r["Artikelbeschrijving"] or r["Artikelnaam"],
                      category=r["Categorie"], min_stock=r["Minimumhoeveelheid"], physical_stock=0, import_batch=batch)
        try:
            res = await db.articles.insert_one(art.to_mongo())
        except DuplicateKeyError:
            skipped.append({"row": r["row"], "article_number": r["Artikelcode"], "reason": "Artikelcode bestaat al"})
            continue
        art.id = str(res.inserted_id)
        if r["Hoeveelheid"] > 0:
            await apply_initial_stock(art, r["Hoeveelheid"], user, f"Beginvoorraad via Excel-import {data.filename}".strip())
        added.append(r["Artikelcode"])
    logger.info("Import %s door %s: %d toegevoegd, %d overgeslagen", batch, user.email, len(added), len(skipped))
    return {"batch": batch, "added_count": len(added), "added": added, "skipped": skipped}


# ---------------------------------------------------------------- SVG-barcode per artikel
MAX_SVG_BYTES = 512 * 1024
SVG_FORBIDDEN_TAGS = {"script", "foreignobject", "iframe", "object", "embed"}


def validate_svg(text: str) -> None:
    low = text.lower()
    if "<!doctype" in low or "<!entity" in low:
        raise HTTPException(400, "SVG-bestand bevat niet-toegestane DOCTYPE/ENTITY-declaraties")
    try:
        root = ET.fromstring(text)
    except ET.ParseError:
        raise HTTPException(400, "Bestand is geen geldig SVG-bestand")
    if root.tag.split("}")[-1].lower() != "svg":
        raise HTTPException(400, "Bestand is geen SVG-afbeelding")
    for el in root.iter():
        if el.tag.split("}")[-1].lower() in SVG_FORBIDDEN_TAGS:
            raise HTTPException(400, "SVG-bestand bevat niet-toegestane elementen (zoals scripts)")
        for k, v in el.attrib.items():
            if k.split("}")[-1].lower().startswith("on") or "javascript:" in str(v).lower():
                raise HTTPException(400, "SVG-bestand bevat niet-toegestane scripts")


class BarcodeSvgIn(BaseModel):
    filename: str = Field(min_length=1, max_length=200)
    svg: str = Field(min_length=1)
    barcode_value: str = Field(min_length=1, max_length=100)


@api.put("/articles/{article_id}/barcode-svg")
async def upload_barcode_svg(article_id: str, data: BarcodeSvgIn, user: User = Depends(require_roles("admin"))):
    if not data.filename.lower().endswith(".svg"):
        raise HTTPException(400, "Alleen SVG-bestanden (.svg) zijn toegestaan")
    raw = data.svg.encode("utf-8")
    if len(raw) > MAX_SVG_BYTES:
        raise HTTPException(400, "SVG-bestand is groter dan 512 KB")
    validate_svg(data.svg)
    value = data.barcode_value.strip()
    art_doc = await db.articles.find_one({"_id": oid(article_id)})
    if not art_doc:
        raise HTTPException(404, "Artikel niet gevonden")
    await ensure_barcodes_free([value], article_id)  # barcodewaarde mag niet bij een ander artikel horen
    path = f"{STORAGE_APP_NAME}/barcodes/{article_id}/{uuid.uuid4().hex}.svg"
    try:
        stored = await run_in_threadpool(put_object, path, raw, "image/svg+xml")
    except HTTPException:
        raise
    except Exception as e:
        logger.error("SVG upload mislukt: %s", e)
        raise HTTPException(502, "Opslaan van het SVG-bestand is mislukt. Probeer het opnieuw.")
    meta = {"path": stored.get("path", path), "filename": data.filename, "size": len(raw), "barcode_value": value,
            "uploaded_at": now(), "uploaded_by": user.id}
    try:
        doc = await db.articles.find_one_and_update(
            {"_id": ObjectId(article_id)},
            {"$set": {"barcode_svg": meta, "updated_at": now()}, "$addToSet": {"barcodes": value}},
            return_document=ReturnDocument.AFTER)
    except DuplicateKeyError:
        raise HTTPException(409, f"Barcode {value} is al gekoppeld aan een ander artikel")
    return (await articles_out([doc]))[0]


@api.get("/articles/{article_id}/barcode-svg")
async def get_barcode_svg(article_id: str, _: User = Depends(current_user)):
    doc = await db.articles.find_one({"_id": oid(article_id)}, {"barcode_svg": 1})
    svg = (doc or {}).get("barcode_svg")
    if not svg:
        raise HTTPException(404, "Geen SVG-barcode gekoppeld")
    try:
        content = await run_in_threadpool(get_object, svg["path"])
    except Exception as e:
        logger.error("SVG ophalen mislukt: %s", e)
        raise HTTPException(502, "SVG-barcode kon niet worden geladen")
    return {"svg": content.decode("utf-8"), "barcode_value": svg["barcode_value"], "filename": svg["filename"]}


@api.delete("/articles/{article_id}/barcode-svg")
async def delete_barcode_svg(article_id: str, _: User = Depends(require_roles("admin"))):
    """Ontkoppelt de SVG-afbeelding. De barcodewaarde blijft bij het artikel (scannen blijft werken)."""
    doc = await db.articles.find_one_and_update({"_id": oid(article_id)}, {"$set": {"barcode_svg": None, "updated_at": now()}},
                                                return_document=ReturnDocument.AFTER)
    if not doc:
        raise HTTPException(404, "Artikel niet gevonden")
    return (await articles_out([doc]))[0]


# ---------------------------------------------------------------- Mutations
def mutation_out(m: Mutation) -> dict:
    d = m.model_dump()
    d["created_at"] = m.created_at.isoformat()
    d.pop("client_request_id", None)
    return d


async def book_mutation(data: MutationIn, user: User) -> dict:
    """Atomische, idempotente boeking.
    1. Mutatie vastleggen met unieke client_request_id (voorkomt dubbele boeking bij herhaald verzenden).
    2. Voorraad wijzigen met één conditionele atomische update (voorkomt negatieve voorraad en race conditions).
    3. Mutatie afronden met voorraad voor/na; bij weigering blijft de poging zichtbaar als 'afgewezen'.
    """
    if user.role not in MUTATION_ROLES[data.type]:
        raise HTTPException(403, "Onvoldoende rechten voor dit mutatietype")
    if data.type != "correctie" and data.quantity <= 0:
        raise HTTPException(422, "Aantal moet groter zijn dan 0")
    art_doc = await db.articles.find_one({"_id": oid(data.article_id)})
    if not art_doc:
        raise HTTPException(404, "Artikel niet gevonden")
    art = Article.from_mongo(art_doc)
    if not art.active:
        raise HTTPException(400, "Artikel is inactief; mutaties zijn niet toegestaan")

    existing = await db.mutations.find_one({"client_request_id": data.client_request_id})
    if existing:
        m = Mutation.from_mongo(existing)
        if m.status == "afgewezen":
            raise HTTPException(409, m.reject_reason or "Mutatie afgewezen")
        return mutation_out(m)

    mut = Mutation(client_request_id=data.client_request_id, article_id=art.id, article_number=art.article_number,
                   article_description=art.title, type=data.type, quantity=data.quantity,
                   user_id=user.id, user_name=user.name, note=data.note.strip())
    try:
        res = await db.mutations.insert_one(mut.to_mongo())
    except DuplicateKeyError:
        return mutation_out(Mutation.from_mongo(await db.mutations.find_one({"client_request_id": data.client_request_id})))
    mid = res.inserted_id
    aid = ObjectId(art.id)

    async def reject(reason: str):
        await db.mutations.update_one({"_id": mid}, {"$set": {"status": "afgewezen", "reject_reason": reason}})
        raise HTTPException(409, reason)

    after = None
    if data.type in ("ontvangst", "retour"):
        delta = data.quantity
        after = await db.articles.find_one_and_update({"_id": aid}, {"$inc": {"physical_stock": delta}, "$set": {"updated_at": now()}},
                                                      return_document=ReturnDocument.AFTER)
    elif data.type == "afboeken":
        delta = -data.quantity
        after = await db.articles.find_one_and_update({"_id": aid, "physical_stock": {"$gte": data.quantity}},
                                                      {"$inc": {"physical_stock": delta}, "$set": {"updated_at": now()}},
                                                      return_document=ReturnDocument.AFTER)
        if not after:
            cur = await db.articles.find_one({"_id": aid}, {"physical_stock": 1})
            await reject(f"Onvoldoende voorraad: {cur['physical_stock']} {art.unit} beschikbaar, {data.quantity} gevraagd")
    else:  # correctie: getelde voorraad wordt nieuwe stand; optimistische lock op huidige stand
        for _ in range(5):
            cur = await db.articles.find_one({"_id": aid}, {"physical_stock": 1})
            delta = data.quantity - cur["physical_stock"]
            after = await db.articles.find_one_and_update({"_id": aid, "physical_stock": cur["physical_stock"]},
                                                          {"$set": {"physical_stock": data.quantity, "updated_at": now()}},
                                                          return_document=ReturnDocument.AFTER)
            if after:
                break
        if not after:
            await reject("Voorraad werd gelijktijdig gewijzigd. Probeer de correctie opnieuw.")

    stock_after = after["physical_stock"]
    doc = await db.mutations.find_one_and_update(
        {"_id": mid},
        {"$set": {"status": "geboekt", "delta": delta, "stock_after": stock_after, "stock_before": stock_after - delta}},
        return_document=ReturnDocument.AFTER)
    return mutation_out(Mutation.from_mongo(doc))


@api.post("/mutations", status_code=201)
async def create_mutation(data: MutationIn, user: User = Depends(current_user)):
    return await book_mutation(data, user)


@api.get("/mutations")
async def list_mutations(
    article_id: str = "",
    type: Optional[MutationType] = None,
    limit: int = Query(100, le=500),
    _: User = Depends(current_user),
):
    filt: dict[str, Any] = {"status": "geboekt"}
    if article_id:
        filt["article_id"] = article_id
    if type:
        filt["type"] = type
    docs = await db.mutations.find(filt).sort("created_at", -1).to_list(limit)
    return [mutation_out(Mutation.from_mongo(d)) for d in docs]


# ---------------------------------------------------------------- Dashboard
@api.get("/dashboard")
async def dashboard(_: User = Depends(current_user)):
    active = await db.articles.count_documents({"active": True})
    below_q = {"active": True, "$expr": {"$lt": ["$physical_stock", "$min_stock"]}}
    below = await db.articles.count_documents(below_q)
    below_docs = await db.articles.find(below_q).sort("physical_stock", 1).to_list(5)
    since = now() - timedelta(days=1)
    today = await db.mutations.count_documents({"status": "geboekt", "created_at": {"$gte": since}})
    recent = await db.mutations.find({"status": "geboekt"}).sort("created_at", -1).to_list(8)
    return {
        "active_articles": active,
        "below_min_count": below,
        "below_min": await articles_out(below_docs),
        "mutations_24h": today,
        "recent_mutations": [mutation_out(Mutation.from_mongo(d)) for d in recent],
    }


# ---------------------------------------------------------------- Startup / seed
@app.on_event("startup")
async def startup():
    await db.users.create_index("email", unique=True)
    await db.articles.create_index("article_number", unique=True)
    await db.articles.create_index("barcodes", unique=True, partialFilterExpression={"barcodes": {"$type": "string"}})
    await db.articles.create_index([("category", 1), ("active", 1)])
    await db.mutations.create_index("client_request_id", unique=True)
    await db.mutations.create_index([("article_id", 1), ("created_at", -1)])
    await db.mutations.create_index([("status", 1), ("created_at", -1)])
    await db.reservations.create_index([("article_id", 1), ("status", 1)])

    seed_users = [
        (os.environ["ADMIN_EMAIL"], os.environ["ADMIN_PASSWORD"], "Beheerder", "admin"),
        ("magazijn@douchebeslag.nl", "Magazijn123!", "Magazijn Medewerker", "magazijn"),
        ("monteur@douchebeslag.nl", "Monteur123!", "Monteur Jan", "monteur"),
    ]
    for email, pw, name, role in seed_users:
        if not await db.users.find_one({"email": email.lower()}):
            await db.users.insert_one(User(email=email.lower(), name=name, role=role, password_hash=hasher.hash(pw)).to_mongo())

    try:
        await run_in_threadpool(init_storage)
    except Exception as e:  # opslag niet bereikbaar: rest van de app blijft werken
        logger.error("Object storage init mislukt: %s", e)


@app.on_event("shutdown")
async def shutdown():
    client.close()


app.include_router(api)
app.add_middleware(CORSMiddleware, allow_credentials=False, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])
