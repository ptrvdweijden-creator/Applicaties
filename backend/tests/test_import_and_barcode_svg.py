"""Tests for new features: Excel import of articles + SVG barcode per article."""
import io
import uuid
import requests
import pytest

BASE = "https://voorraad-scan.preview.emergentagent.com"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------------- Excel template ----------------
class TestImportTemplate:
    def test_template_download_xlsx(self, h_admin):
        r = requests.get(f"{BASE}/api/import/articles/template", headers=h_admin)
        assert r.status_code == 200
        assert XLSX_MIME in r.headers.get("content-type", "")
        # xlsx = zip starting with PK
        assert r.content[:2] == b"PK"
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(r.content))
            ws = wb.active
            headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            assert headers == ["Artikelnaam", "Artikelcode", "Artikelbeschrijving",
                               "Hoeveelheid", "Categorie", "Minimumhoeveelheid"]
        except ImportError:
            pytest.skip("openpyxl not installed in test env")


# ---------------- Preview ----------------
class TestImportPreview:
    def test_monteur_forbidden(self, h_mont):
        with open("/tmp/test_import.xlsx", "rb") as f:
            r = requests.post(f"{BASE}/api/import/articles/preview", headers=h_mont,
                              files={"file": ("test_import.xlsx", f, XLSX_MIME)})
        assert r.status_code == 403

    def test_non_xlsx_rejected(self, h_admin):
        r = requests.post(f"{BASE}/api/import/articles/preview", headers=h_admin,
                          files={"file": ("foo.csv", b"a,b\n1,2\n", "text/csv")})
        assert r.status_code == 400

    def test_missing_column_rejected(self, h_admin):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        # missing "Hoeveelheid"
        ws.append(["Artikelnaam", "Artikelcode", "Artikelbeschrijving",
                   "Categorie", "Minimumhoeveelheid"])
        ws.append(["Naam", "TEST-MC-1", "", "Chroom", 1])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        r = requests.post(f"{BASE}/api/import/articles/preview", headers=h_admin,
                          files={"file": ("bad.xlsx", buf.getvalue(), XLSX_MIME)})
        assert r.status_code == 400
        assert "Hoeveelheid" in r.text or "kolom" in r.text.lower()

    def test_preview_valid_and_invalid_rows(self, h_admin):
        with open("/tmp/test_import.xlsx", "rb") as f:
            r = requests.post(f"{BASE}/api/import/articles/preview", headers=h_admin,
                              files={"file": ("test_import.xlsx", f, XLSX_MIME)})
        assert r.status_code == 200, r.text
        data = r.json()
        # 4 data rows in /tmp/test_import.xlsx
        assert len(data["rows"]) == 4
        rows = {r["row"]: r for r in data["rows"]}
        # row 2 TEST-X-001 valid (unless already exists)
        # row 3 TEST-X-002 valid
        assert rows[3]["valid"] is True
        # row 4: duplicate code of row 2 in file -> error
        row4 = rows[4]
        assert row4["valid"] is False
        assert any("Dubbele" in e or "bestaat" in e.lower() for e in row4["errors"])
        # row 5: Artikelnaam ontbreekt
        row5 = rows[5]
        assert row5["valid"] is False
        assert any("Artikelnaam" in e for e in row5["errors"])
        assert data["valid_count"] >= 1
        assert data["invalid_count"] >= 2


# ---------------- Commit ----------------
class TestImportCommit:
    def test_commit_and_persistence(self, h_admin):
        # Create a unique xlsx so repeated test runs don't collide
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["Artikelnaam", "Artikelcode", "Artikelbeschrijving",
                   "Hoeveelheid", "Categorie", "Minimumhoeveelheid"])
        suffix = uuid.uuid4().hex[:6].upper()
        code_a = f"TEST-IMP-{suffix}-A"
        code_b = f"TEST-IMP-{suffix}-B"
        ws.append([f"TEST import {suffix} A", code_a, "beschr A", 12, "TestCat", 4])
        ws.append([f"TEST import {suffix} B", code_b, "", 3, "TestCat", 1])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        # preview
        pr = requests.post(f"{BASE}/api/import/articles/preview", headers=h_admin,
                          files={"file": (f"imp_{suffix}.xlsx", buf.getvalue(), XLSX_MIME)})
        assert pr.status_code == 200, pr.text
        pv = pr.json()
        assert pv["valid_count"] == 2
        assert pv["invalid_count"] == 0

        payload = {"filename": pv["filename"], "rows": pv["rows"]}
        cr = requests.post(f"{BASE}/api/import/articles/commit", headers=h_admin, json=payload)
        assert cr.status_code == 200, cr.text
        res = cr.json()
        assert res["added_count"] == 2
        assert sorted(res["added"]) == sorted([code_a, code_b])
        assert res["skipped"] == []

        # Second commit (same payload) -> all skipped as existing, nothing added
        cr2 = requests.post(f"{BASE}/api/import/articles/commit", headers=h_admin, json=payload)
        assert cr2.status_code == 200
        res2 = cr2.json()
        assert res2["added_count"] == 0
        assert len(res2["skipped"]) == 2
        for s in res2["skipped"]:
            assert "bestaat" in s["reason"].lower() or "al" in s["reason"].lower()

        # Verify article A: physical_stock = 12 and exactly one beginvoorraad mutation (no ontvangst)
        arts = requests.get(f"{BASE}/api/articles", headers=h_admin,
                            params={"q": code_a}).json()
        assert arts, f"article {code_a} not found"
        art_a = next(a for a in arts if a["article_number"] == code_a)
        assert art_a["physical_stock"] == 12
        aid = art_a["id"]
        muts = requests.get(f"{BASE}/api/mutations", headers=h_admin,
                            params={"article_id": aid}).json()
        begin = [m for m in muts if m["type"] == "beginvoorraad"]
        ontv = [m for m in muts if m["type"] == "ontvangst"]
        assert len(begin) == 1
        assert begin[0]["quantity"] == 12
        assert ontv == []

        # Afboeken + ontvangst still work on imported article
        afb = requests.post(f"{BASE}/api/mutations", headers=h_admin, json={
            "client_request_id": f"test-{uuid.uuid4()}",
            "article_id": aid, "type": "afboeken", "quantity": 2, "note": "TEST"
        })
        assert afb.status_code == 201, afb.text
        assert afb.json()["stock_after"] == 10
        ont = requests.post(f"{BASE}/api/mutations", headers=h_admin, json={
            "client_request_id": f"test-{uuid.uuid4()}",
            "article_id": aid, "type": "ontvangst", "quantity": 5, "note": "TEST"
        })
        assert ont.status_code == 201
        assert ont.json()["stock_after"] == 15


# ---------------- SVG barcode ----------------
def _mk_article(h_admin):
    code = f"TEST-SVG-{uuid.uuid4().hex[:6].upper()}"
    r = requests.post(f"{BASE}/api/articles", headers=h_admin, json={
        "article_number": code, "name": code, "description": f"TEST svg {code}",
        "category": "TestCat", "min_stock": 1
    })
    # some servers may not accept 'name', fallback
    if r.status_code >= 400:
        r = requests.post(f"{BASE}/api/articles", headers=h_admin, json={
            "article_number": code, "description": f"TEST svg {code}",
            "category": "TestCat", "min_stock": 1
        })
    assert r.status_code == 201, r.text
    return r.json(), code


SVG_OK = open("/tmp/test_barcode.svg").read()


class TestBarcodeSvg:
    def test_rejects_non_svg_filename(self, h_admin):
        art, _ = _mk_article(h_admin)
        r = requests.put(f"{BASE}/api/articles/{art['id']}/barcode-svg", headers=h_admin,
                        json={"filename": "foo.png", "svg": SVG_OK,
                              "barcode_value": f"TBV-{uuid.uuid4().hex[:8]}"})
        assert r.status_code == 400

    def test_rejects_script(self, h_admin):
        art, _ = _mk_article(h_admin)
        bad = '<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>'
        r = requests.put(f"{BASE}/api/articles/{art['id']}/barcode-svg", headers=h_admin,
                        json={"filename": "x.svg", "svg": bad,
                              "barcode_value": f"TBV-{uuid.uuid4().hex[:8]}"})
        assert r.status_code == 400

    def test_upload_and_get_and_delete(self, h_admin):
        art, _ = _mk_article(h_admin)
        value = f"TBV-{uuid.uuid4().hex[:10].upper()}"
        r = requests.put(f"{BASE}/api/articles/{art['id']}/barcode-svg", headers=h_admin,
                        json={"filename": "code.svg", "svg": SVG_OK, "barcode_value": value})
        assert r.status_code == 200, r.text
        updated = r.json()
        assert value in updated["barcodes"]

        # GET returns svg + value
        gr = requests.get(f"{BASE}/api/articles/{art['id']}/barcode-svg", headers=h_admin)
        assert gr.status_code == 200
        gd = gr.json()
        assert "<svg" in gd["svg"]
        assert gd["barcode_value"] == value
        assert gd["filename"] == "code.svg"

        # by-barcode returns correct article
        br = requests.get(f"{BASE}/api/articles/by-barcode/{value}", headers=h_admin)
        assert br.status_code == 200
        assert br.json()["id"] == art["id"]

        # DELETE unlinks svg but keeps barcode value
        dr = requests.delete(f"{BASE}/api/articles/{art['id']}/barcode-svg", headers=h_admin)
        assert dr.status_code == 200
        assert value in dr.json()["barcodes"]
        # GET should now 404
        gr2 = requests.get(f"{BASE}/api/articles/{art['id']}/barcode-svg", headers=h_admin)
        assert gr2.status_code == 404
        # by-barcode still works (value kept)
        br2 = requests.get(f"{BASE}/api/articles/by-barcode/{value}", headers=h_admin)
        assert br2.status_code == 200
        assert br2.json()["id"] == art["id"]

    def test_duplicate_barcode_value_rejected(self, h_admin):
        art1, _ = _mk_article(h_admin)
        art2, _ = _mk_article(h_admin)
        value = f"TBV-{uuid.uuid4().hex[:10].upper()}"
        r1 = requests.put(f"{BASE}/api/articles/{art1['id']}/barcode-svg", headers=h_admin,
                        json={"filename": "a.svg", "svg": SVG_OK, "barcode_value": value})
        assert r1.status_code == 200
        r2 = requests.put(f"{BASE}/api/articles/{art2['id']}/barcode-svg", headers=h_admin,
                        json={"filename": "b.svg", "svg": SVG_OK, "barcode_value": value})
        assert r2.status_code == 409

    def test_patch_cannot_remove_svg_linked_barcode(self, h_admin):
        art, _ = _mk_article(h_admin)
        value = f"TBV-{uuid.uuid4().hex[:10].upper()}"
        r = requests.put(f"{BASE}/api/articles/{art['id']}/barcode-svg", headers=h_admin,
                       json={"filename": "a.svg", "svg": SVG_OK, "barcode_value": value})
        assert r.status_code == 200
        # Try to PATCH article removing the barcode (set to empty list)
        full = r.json()
        other = [b for b in full["barcodes"] if b != value]
        pr = requests.patch(f"{BASE}/api/articles/{art['id']}", headers=h_admin,
                           json={"description": full["description"], "barcodes": other})
        assert pr.status_code == 409

    def test_monteur_cannot_upload_svg(self, h_admin, h_mont):
        art, _ = _mk_article(h_admin)
        r = requests.put(f"{BASE}/api/articles/{art['id']}/barcode-svg", headers=h_mont,
                        json={"filename": "a.svg", "svg": SVG_OK,
                              "barcode_value": f"TBV-{uuid.uuid4().hex[:8]}"})
        assert r.status_code == 403
